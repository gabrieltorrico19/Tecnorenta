import io
from datetime import date

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.activo import Activo, EstadoActivo

router = APIRouter(prefix="/exportar", tags=["Exportar"])

# US-20: Exportar inventario filtrado a Excel/PDF
@router.get("/inventario/excel")
def exportar_inventario_excel(
    estado: str = Query(None, description="Filtrar por estado: disponible, rentado, mantenimiento, baja"),
    id_categoria: int = Query(None, description="Filtrar por categoría"),
    db: Session = Depends(get_db),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        from fastapi import HTTPException
        raise HTTPException(
            status_code=500,
            detail="Dependencia 'openpyxl' no instalada. Ejecutar: pip install openpyxl",
        )

    query = db.query(Activo)
    if estado:
        query = query.filter(Activo.estado == estado)
    if id_categoria:
        query = query.filter(Activo.id_categoria == id_categoria)
    activos = query.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario de Activos"

    # Encabezado con estilo
    encabezados = ["ID", "Código Inventario", "Modelo", "N° Serie", "Estado", "Fecha Compra", "Valor Depreciado (BOB)", "Categoría"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col, titulo in enumerate(encabezados, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Datos
    for row, activo in enumerate(activos, start=2):
        ws.cell(row=row, column=1, value=activo.id)
        ws.cell(row=row, column=2, value=activo.codigo_inventario)
        ws.cell(row=row, column=3, value=activo.modelo)
        ws.cell(row=row, column=4, value=activo.numero_serie)
        ws.cell(row=row, column=5, value=activo.estado.value if activo.estado else "")
        ws.cell(row=row, column=6, value=str(activo.fecha_compra) if activo.fecha_compra else "")
        ws.cell(row=row, column=7, value=activo.valor_depreciado)
        ws.cell(row=row, column=8, value=str(activo.id_categoria) if activo.id_categoria else "")

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    # Fila de resumen
    ws.append([])
    ws.append(["Total de activos exportados:", len(activos), "", "", f"Generado: {date.today()}"])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"inventario_activos_{date.today()}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
